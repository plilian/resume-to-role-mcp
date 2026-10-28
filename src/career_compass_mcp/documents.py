from __future__ import annotations

import re
import unicodedata
from pathlib import Path

from docx import Document
from docx.shared import Inches,Pt
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle,getSampleStyleSheet
from reportlab.platypus import Paragraph,SimpleDocTemplate,Spacer
from reportlab.lib import colors
from xml.sax.saxutils import escape

def safe_slug(text: str) -> str:
    text=unicodedata.normalize("NFC",str(text or ""))
    slug=re.sub(r"[^\w.-]+","_",text,flags=re.UNICODE).strip("_ .")[:70]
    if slug.upper() in {"CON","PRN","AUX","NUL",*(f"COM{i}" for i in range(1,10)),*(f"LPT{i}" for i in range(1,10))}:
        slug="_"+slug
    return slug or "Role"

def split_resume(path: str) -> tuple[str,list[str]]:
    p=Path(path).expanduser().resolve(strict=True)
    if not p.is_file() or p.suffix.lower() not in {".pdf",".docx",".txt",".md"}:
        raise ValueError(f"Unsupported resume file: {p.name}")
    if p.stat().st_size>20*1024*1024: raise ValueError(f"File is larger than 20 MB: {p.name}")
    if p.suffix.lower()==".pdf":
        from pypdf import PdfReader
        text="\n".join(page.extract_text() or "" for page in PdfReader(str(p)).pages)
    elif p.suffix.lower()==".docx":
        d=Document(str(p)); text="\n".join([x.text for x in d.paragraphs]+[" | ".join(c.text for c in row.cells) for t in d.tables for row in t.rows])
    else: text=p.read_text(encoding="utf-8",errors="replace")
    text=re.sub(r"[ \t]+"," ",text); text=re.sub(r"\n{3,}","\n\n",text).strip()
    if not text: raise ValueError(f"No readable text found in {p.name}; scanned-image PDFs need OCR.")
    chunks=[x.strip() for x in re.split(r"\n\s*\n|(?<=[.!?])\s+(?=[A-Z0-9])",text) if len(x.strip())>18]
    return text[:300000],chunks[:500]

def _font_path(language: str) -> str | None:
    candidates=[]
    if __import__("platform").system()=="Windows":
        candidates.extend([r"C:\Windows\Fonts\arial.ttf",r"C:\Windows\Fonts\tahoma.ttf"])
    elif __import__("platform").system()=="Darwin":
        candidates.extend(["/System/Library/Fonts/Supplemental/Arial.ttf","/System/Library/Fonts/Supplemental/Arial Unicode.ttf"])
    candidates.extend([
        "/usr/share/fonts/truetype/noto/NotoSansArabic-Regular.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ])
    return next((path for path in candidates if Path(path).is_file()),None)


def _set_rtl_paragraph(paragraph, language: str) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    ppr=paragraph._p.get_or_add_pPr()
    bidi=OxmlElement("w:bidi"); bidi.set(qn("w:val"),"1"); ppr.append(bidi)
    paragraph.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    locale={"fa":"fa-IR","ar":"ar-SA"}.get(language,"ar-SA")
    for run in paragraph.runs:
        rpr=run._r.get_or_add_rPr()
        rtl=OxmlElement("w:rtl"); rtl.set(qn("w:val"),"1"); rpr.append(rtl)
        lang=OxmlElement("w:lang"); lang.set(qn("w:val"),locale); rpr.append(lang)


def render_docs(content: str, kind: str, name: str, title: str, out_dir: Path,
                language: str = "en") -> tuple[str,str]:
    out_dir.mkdir(parents=True,exist_ok=True)
    stem=f"{safe_slug(name or 'Candidate')}_{safe_slug(title)}_{__import__('datetime').date.today().year}"
    docx=out_dir/(stem+".docx"); pdf=out_dir/(stem+".pdf")
    rtl=language in {"fa","ar"}
    d=Document(); s=d.sections[0]; s.top_margin=Inches(.55); s.bottom_margin=Inches(.55); s.left_margin=Inches(.65); s.right_margin=Inches(.65)
    d.styles["Normal"].font.name="Arial"; d.styles["Normal"].font.size=Pt(9.5)
    for line in content.splitlines():
        t=line.strip()
        if not t: paragraph=d.add_paragraph()
        elif t.startswith("#"):
            paragraph=d.add_paragraph(); run=paragraph.add_run(t.lstrip("# ").strip()); run.bold=True; run.font.size=Pt(13)
        elif t.startswith(("- ","* ","• ")): paragraph=d.add_paragraph(t[2:].strip(),style="List Bullet")
        elif len(t)<80 and t.isupper(): paragraph=d.add_paragraph(); run=paragraph.add_run(t); run.bold=True
        else: paragraph=d.add_paragraph(t)
        if rtl: _set_rtl_paragraph(paragraph,language)
    d.save(docx)

    styles=getSampleStyleSheet(); rtl_font=_font_path(language)
    if rtl and not rtl_font:
        raise RuntimeError("No Arabic/Persian-capable system font was found. Install Noto Sans Arabic, Arial, or Tahoma, then retry PDF export.")
    font_name="Helvetica"
    if rtl_font:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont as ReportLabFont
        font_name="CareerCompassUnicode"
        if font_name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(ReportLabFont(font_name,rtl_font))
    body=ParagraphStyle("CareerCompassBody",parent=styles["BodyText"],fontName=font_name,fontSize=9,leading=12,spaceAfter=4,textColor=colors.HexColor("#202124"),alignment=2 if rtl else 0)
    heading=ParagraphStyle("CareerCompassHeading",parent=styles["Heading2"],fontName=font_name,fontSize=13,leading=16,spaceAfter=5,alignment=2 if rtl else 0)
    if rtl:
        import arabic_reshaper
        from bidi.algorithm import get_display
        transform=lambda value:get_display(arabic_reshaper.reshape(value))
    else:
        transform=lambda value:value
    story=[]
    for line in content.splitlines():
        t=line.strip()
        if not t: story.append(Spacer(1,5)); continue
        if t.startswith("#"):
            story.append(Paragraph("<b>"+escape(transform(t.lstrip("# ").strip()))+"</b>",heading)); continue
        if t.startswith(("- ","* ","• ")): t="• "+t[2:].strip()
        story.append(Paragraph(escape(transform(t)),body))
    SimpleDocTemplate(str(pdf),pagesize=letter,rightMargin=42,leftMargin=42,topMargin=36,bottomMargin=36).build(story)
    return str(docx),str(pdf)

def numbers(text: str) -> set[str]:
    digits="۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩"
    text=str(text).translate(str.maketrans(digits,"01234567890123456789")).replace("٪","%").replace("٪","%")
    return {re.sub(r"[, ]","",x) for x in re.findall(r"(?<!\w)\d[\d,.%+/-]*(?!\w)",text,flags=re.UNICODE)}
